"""Biomarker dictionary: allowed combinations biomarker x result type (mutation/variant, expression, fusion, ...).

Source of truth is trialink/data/biomarkery.csv (converted from the clinical overview biomarkery_prehled_upraveno.xlsx),
imported into IRIS table TriaLink.BiomarkerResultType. Codes are the marker names used by study criteria and matching.
Invalid combinations only produce warnings – the entered value is still used for matching.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterable

from .model import BiomarkerResult

DATA_FILE = Path(__file__).resolve().parent / "data" / "biomarkery.csv"
RESULT_TYPES = {
    "mutace": "Mutace / varianta",
    "exprese": "Exprese / IHC",
    "fuze": "Fúze / přestavba",
    "amplifikace": "Amplifikace / počet kopií",
    "delece": "Delece / ztráta",
    "metylace": "Metylace",
    "jiny": "Jiný typ výsledku",
}
COLUMNS = ["kod", "nazev", "opat_kod", *RESULT_TYPES, "zdroj"]
SOURCE_EXCEL = "biomarkery_prehled_upraveno.xlsx"
SOURCE_CRITERIA = "doplněno z kritérií studií"
# Markers used by study criteria but missing in the clinical overview (agreed supplement).
SUPPLEMENT = [
    ("CCNE1", "CCNE1 amplifikace", {"amplifikace": ""}),
    ("MTAP", "MTAP ztráta / exprese", {"delece": "", "exprese": ""}),
    ("HPV", "HPV", {"jiny": "Virová pozitivita"}),
    ("TROP2", "TROP2", {"exprese": ""}),
]
# Excel/OPAT labels whose code cannot be derived by dropping the result-type word (keys are normalized labels).
CODE_OVERRIDES = {
    "pdl1mutace": "PD-L1 mutace", "kodelecia1p19q": "1p/19q", "kodelece1p19q": "1p/19q",
    "prestavbaret": "RET", "proliferaciaki67": "Ki67", "proliferaceki67": "Ki67", "stromalnitils": "TILs",
    "gishrd": "GIS/HRD", "atrxmutztrexprese": "ATRX", "cldn182": "CLDN18.2", "esr1": "ESR1",
}
_TYPE_WORDS = r"(mutace|mutacia|amplifikace|amplifikacia|fuzia|fuze|alterace|metylace|metylacia)"


@dataclass(frozen=True)
class BiomarkerTypeRow:
    code: str
    label: str
    opat_kod: int | None
    result_type: str
    note: str
    source: str


def _key(text: str | None) -> str:
    """Case/diacritics/punctuation-insensitive key: 'Ki-67' == 'KI67', 'CLDN 18-2' == 'CLDN18.2'."""
    text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", text.lower())


def derive_code(label: str) -> str:
    label = re.sub(r"\s+", " ", label).strip()
    if _key(label) in CODE_OVERRIDES:
        return CODE_OVERRIDES[_key(label)]
    ascii_label = unicodedata.normalize("NFKD", label).encode("ascii", "ignore").decode("ascii")
    stripped = re.sub(rf"^(exprese|prestavba|proliferace|proliferacia)\s+|\s+{_TYPE_WORDS}$", "", ascii_label, flags=re.I)
    if stripped == ascii_label:
        return label
    # Keep the original (accented) spelling of the remaining part.
    start = ascii_label.find(stripped)
    return label[start:start + len(stripped)]


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1250")


def parse(data: bytes) -> tuple[list[BiomarkerTypeRow], list[str]]:
    """Wide CSV (one line per biomarker, mark or note per result type) -> long rows. Invalid lines are reported."""
    reader = csv.DictReader(io.StringIO(_decode(data)), delimiter=";")
    if reader.fieldnames != COLUMNS:
        raise ValueError(f"Neočekávané sloupce číselníku biomarkerů: {reader.fieldnames}")
    rows, errors, seen = [], [], {}
    for n, r in enumerate(reader, start=2):
        code, label = (r["kod"] or "").strip(), (r["nazev"] or "").strip()
        if not code or not label:
            errors.append(f"řádek {n}: chybí kód nebo název")
            continue
        if _key(code) in seen:
            errors.append(f"řádek {n}: duplicitní kód {code} (řádek {seen[_key(code)]})")
            continue
        seen[_key(code)] = n
        opat = (r["opat_kod"] or "").strip()
        if opat and not opat.isdigit():
            errors.append(f"řádek {n}: neplatný opat_kod {opat!r}")
            continue
        types = [(t, (r[t] or "").strip()) for t in RESULT_TYPES if (r[t] or "").strip()]
        if not types:
            errors.append(f"řádek {n}: {code} nemá žádný povolený typ výsledku")
            continue
        for t, mark in types:
            note = "" if mark.lower() in ("x", "✓", "1", "ano") else mark
            rows.append(BiomarkerTypeRow(code, label, int(opat) if opat else None, t, note, (r["zdroj"] or "").strip()))
    return rows, errors


def to_csv(rows: Iterable[BiomarkerTypeRow]) -> str:
    by_code: dict[str, dict] = {}
    for r in rows:
        line = by_code.setdefault(r.code, {"kod": r.code, "nazev": r.label, "opat_kod": r.opat_kod or "", "zdroj": r.source,
                                           **{t: "" for t in RESULT_TYPES}})
        line[r.result_type] = r.note or "x"
    out = io.StringIO()
    writer = csv.DictWriter(out, COLUMNS, delimiter=";", lineterminator="\n")
    writer.writeheader()
    writer.writerows(by_code.values())
    return out.getvalue()


def _opat_codes(path: Path | None) -> dict[str, int]:
    """Biomarker code -> OPAT 'prediktor' code from the customer dictionary (optional)."""
    if not path or not path.is_file():
        return {}
    import openpyxl

    ws = openpyxl.load_workbook(path, read_only=True)["prediktor"]
    return {_key(derive_code(str(label))): int(kod) for kod, label in ws.iter_rows(min_row=2, values_only=True) if label and kod}


def from_xlsx(path: Path, opat_path: Path | None = None) -> tuple[list[BiomarkerTypeRow], list[str]]:
    """Convert the clinical overview Excel (first sheet: Biomarker + 7 result-type columns, ✓ or a note)."""
    import openpyxl

    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    lines = list(ws.iter_rows(values_only=True))
    header = [str(c or "").strip() for c in lines[0][:8]]
    if header[0] != "Biomarker" or len(header) < 8:
        raise ValueError(f"Neočekávané záhlaví listu {ws.title!r}: {header}")
    opat = _opat_codes(opat_path)
    rows, errors = [], []
    for n, line in enumerate(lines[1:], start=2):
        label = str(line[0] or "").strip()
        if not label or label == "Hodnoty":
            continue
        marks = {t: str(v).strip() for t, v in zip(RESULT_TYPES, line[1:8]) if v not in (None, "")}
        if not marks:
            errors.append(f"řádek {n}: {label} nemá žádný typ výsledku")
            continue
        code = derive_code(label)
        variants = [(code, label, marks)]
        if code == "PD-L1":  # criteria and reports distinguish the TPS and CPS score
            base = {t: m for t, m in marks.items() if t != "jiny"}
            variants = [(f"PD-L1 {s}", f"{label} – {s}", {**base, "jiny": f"{s} skóre"}) for s in ("TPS", "CPS")]
        for c, lab, m in variants:
            for t, mark in m.items():
                note = "" if mark in ("✓", "x", "X") else mark
                rows.append(BiomarkerTypeRow(c, lab, opat.get(_key(code)), t, note, SOURCE_EXCEL))
    known = {_key(r.code) for r in rows}
    for code, label, marks in SUPPLEMENT:
        if _key(code) not in known:
            rows += [BiomarkerTypeRow(code, label, None, t, note, SOURCE_CRITERIA) for t, note in marks.items()]
    return rows, errors


class BiomarkerDictionary:
    def __init__(self, rows: Iterable[BiomarkerTypeRow]):
        self.types: dict[str, dict[str, str]] = {}
        self.labels: dict[str, str] = {}
        self._by_key: dict[str, str] = {}
        for r in rows:
            self.types.setdefault(r.code, {})[r.result_type] = r.note
            self.labels[r.code] = r.label
        for code, label in self.labels.items():
            self._by_key.setdefault(_key(label), code)
            for part in code.split("/"):
                if len(_key(part)) >= 3:
                    self._by_key.setdefault(_key(part), code)
        for code in self.labels:  # exact codes win over label/part aliases
            self._by_key[_key(code)] = code

    @property
    def codes(self) -> list[str]:
        return sorted(self.labels, key=str.casefold)

    def canonical(self, marker: str | None) -> str | None:
        return self._by_key.get(_key(marker))

    def label(self, code: str) -> str:
        label = self.labels.get(code, code)
        return code if _key(label) == _key(code) else f"{code} – {label}"

    def allowed(self, code: str) -> list[str]:
        return [t for t in RESULT_TYPES if t in self.types.get(code, {})]

    def describe(self, code: str) -> str:
        return ", ".join(RESULT_TYPES[t] + (f" ({self.types[code][t]})" if self.types[code][t] else "") for t in self.allowed(code))

    def check(self, b: BiomarkerResult) -> list[str]:
        """Czech warnings for combinations biomarker x result type x variant x value not allowed by the dictionary."""
        code = self.canonical(b.marker)
        if code is None:
            return [f"{b.marker}: biomarker není v číselníku – zkontrolujte název (typ výsledku a variantu nelze ověřit)."]
        allowed, issues = self.allowed(code), []
        names = ", ".join(RESULT_TYPES[t] for t in allowed)
        if b.result_type and b.result_type not in allowed:
            issues.append(f"{code}: typ „{RESULT_TYPES.get(b.result_type, b.result_type)}“ není pro tento biomarker povolen "
                          f"(povoleno: {names}).")
        if b.variant:
            if "mutace" not in allowed:
                issues.append(f"{code}: varianta se u tohoto biomarkeru nezadává (povoleno: {names}).")
            elif b.result_type and b.result_type != "mutace":
                issues.append(f"{code}: varianta patří jen k typu „{RESULT_TYPES['mutace']}“.")
        if b.values and not {"exprese", "jiny"} & set(allowed):
            issues.append(f"{code}: hodnota (% / skóre) se u tohoto biomarkeru nezadává (povoleno: {names}).")
        return issues

    def check_all(self, results: Iterable[BiomarkerResult]) -> list[str]:
        return [issue for b in results for issue in self.check(b)]


def load_file(path: Path | None = None) -> tuple[list[BiomarkerTypeRow], list[str], str, str]:
    path = path or DATA_FILE
    data = path.read_bytes()
    rows, errors = parse(data)
    return rows, errors, path.name, hashlib.sha256(data).hexdigest()


@lru_cache(maxsize=1)
def default_dictionary() -> BiomarkerDictionary:
    """Packaged CSV – used outside Streamlit (extraction, prompts) and as fallback without IRIS."""
    return BiomarkerDictionary(load_file()[0])


if __name__ == "__main__":
    # python -m trialink.biomarkers --from-xlsx biomarkery_prehled_upraveno.xlsx [--opat <OPAT xlsx>]  – regenerate the CSV
    # python -m trialink.biomarkers [--check]                                                           – validate / import into IRIS
    import sys

    from dotenv import load_dotenv

    args = sys.argv[1:]
    if "--from-xlsx" in args:
        src = Path(args[args.index("--from-xlsx") + 1])
        opat = Path(args[args.index("--opat") + 1]) if "--opat" in args else next(
            Path(__file__).resolve().parent.parent.glob("data/*Biomarkery/Biomarkery_OPAT_ciselniky_*.xlsx"), None)
        rows, errors = from_xlsx(src, opat)
        for e in errors:
            print("  CHYBA", e)
        if errors:
            raise SystemExit(1)
        DATA_FILE.write_text(to_csv(rows), encoding="utf-8", newline="")
        print(f"{DATA_FILE.name}: {len({r.code for r in rows})} biomarkerů, {len(rows)} kombinací (OPAT kódy: {opat.name if opat else 'ne'})")
    rows, errors, name, digest = load_file()
    print(f"{name}: {len({r.code for r in rows})} biomarkerů, {len(rows)} kombinací, chyb {len(errors)}, sha256={digest[:12]}…")
    for e in errors[:20]:
        print("  CHYBA", e)
    if errors or "--check" in args or "--from-xlsx" in args:
        raise SystemExit(1 if errors else 0)
    load_dotenv()
    from . import db

    conn = db.connect()
    try:
        db.ensure_schema(conn)
        print("IRIS:", db.save_biomarker_types(conn, rows, name, digest))
    finally:
        conn.close()
